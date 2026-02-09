"""
Create test documents with footnotes for comparison testing.
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_original_with_footnotes():
    """Create an original document with footnotes."""
    doc = Document()

    # Title
    title = doc.add_heading('Service Agreement', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Introduction with footnote reference
    para1 = doc.add_paragraph()
    para1.add_run('This Service Agreement ("Agreement") is entered into as of January 1, 2026')
    # Note: python-docx doesn't have built-in footnote support, so we'll simulate with superscript
    run = para1.add_run('1')
    run.font.superscript = True
    para1.add_run(' between Valon Technologies LLC ("Provider") and the Client ("Client").')

    # Section 1
    doc.add_heading('1. Services', level=1)
    para2 = doc.add_paragraph()
    para2.add_run('Provider agrees to provide mortgage servicing technology services')
    run = para2.add_run('2')
    run.font.superscript = True
    para2.add_run(' as described in Exhibit A attached hereto.')

    # Section 2
    doc.add_heading('2. Term', level=1)
    para3 = doc.add_paragraph()
    para3.add_run('The initial term of this Agreement shall be one (1) year')
    run = para3.add_run('3')
    run.font.superscript = True
    para3.add_run(' commencing on the Effective Date.')

    # Section 3
    doc.add_heading('3. Compensation', level=1)
    para4 = doc.add_paragraph()
    para4.add_run('Client shall pay Provider a monthly fee of $10,000')
    run = para4.add_run('4')
    run.font.superscript = True
    para4.add_run(' for the services described herein.')

    # Section 4
    doc.add_heading('4. Confidentiality', level=1)
    para5 = doc.add_paragraph(
        'Each party agrees to maintain the confidentiality of all proprietary information '
        'disclosed by the other party during the term of this Agreement.'
    )

    # Footnotes section
    doc.add_page_break()
    doc.add_heading('Footnotes', level=1)
    doc.add_paragraph('1. Subject to regulatory approval.')
    doc.add_paragraph('2. As defined in Section 1.1 of the Master Agreement.')
    doc.add_paragraph('3. With automatic renewal for successive one-year terms.')
    doc.add_paragraph('4. Payment due within 30 days of invoice.')

    return doc


def create_modified_with_footnotes():
    """Create a modified document with footnotes (with changes)."""
    doc = Document()

    # Title
    title = doc.add_heading('Service Agreement', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Introduction with footnote reference - CHANGED DATE
    para1 = doc.add_paragraph()
    para1.add_run('This Service Agreement ("Agreement") is entered into as of March 15, 2026')
    run = para1.add_run('1')
    run.font.superscript = True
    para1.add_run(' between Valon Technologies LLC ("Provider") and the Client ("Client").')

    # Section 1 - ADDED more detail
    doc.add_heading('1. Services', level=1)
    para2 = doc.add_paragraph()
    para2.add_run('Provider agrees to provide comprehensive mortgage servicing technology services')
    run = para2.add_run('2')
    run.font.superscript = True
    para2.add_run(' including software licensing and support as described in Exhibit A attached hereto.')

    # Section 2 - CHANGED term
    doc.add_heading('2. Term', level=1)
    para3 = doc.add_paragraph()
    para3.add_run('The initial term of this Agreement shall be two (2) years')
    run = para3.add_run('3')
    run.font.superscript = True
    para3.add_run(' commencing on the Effective Date.')

    # Section 3 - CHANGED amount
    doc.add_heading('3. Compensation', level=1)
    para4 = doc.add_paragraph()
    para4.add_run('Client shall pay Provider a monthly fee of $15,000')
    run = para4.add_run('4')
    run.font.superscript = True
    para4.add_run(' for the services described herein, with annual increases of 3%.')

    # Section 4 - UNCHANGED
    doc.add_heading('4. Confidentiality', level=1)
    para5 = doc.add_paragraph(
        'Each party agrees to maintain the confidentiality of all proprietary information '
        'disclosed by the other party during the term of this Agreement.'
    )

    # NEW Section 5
    doc.add_heading('5. Limitation of Liability', level=1)
    para6 = doc.add_paragraph()
    para6.add_run('Neither party shall be liable for indirect, incidental, or consequential damages')
    run = para6.add_run('5')
    run.font.superscript = True
    para6.add_run(' arising out of this Agreement.')

    # Footnotes section - MODIFIED
    doc.add_page_break()
    doc.add_heading('Footnotes', level=1)
    doc.add_paragraph('1. Subject to regulatory approval and board authorization.')
    doc.add_paragraph('2. As defined in Section 1.1 of the Master Agreement and applicable schedules.')
    doc.add_paragraph('3. With automatic renewal for successive two-year terms unless terminated.')
    doc.add_paragraph('4. Payment due within 45 days of invoice. Late payments subject to 1.5% monthly interest.')
    doc.add_paragraph('5. Except in cases of gross negligence or willful misconduct.')

    return doc


def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))

    # Create original document
    original = create_original_with_footnotes()
    original_path = os.path.join(script_dir, 'footnote_test_v1.docx')
    original.save(original_path)
    print(f"Created: {original_path}")

    # Create modified document
    modified = create_modified_with_footnotes()
    modified_path = os.path.join(script_dir, 'footnote_test_v2.docx')
    modified.save(modified_path)
    print(f"Created: {modified_path}")

    print("\nFootnote test documents created successfully!")
    print("\nKey differences:")
    print("- Date changed: January 1 -> March 15")
    print("- Services: Added 'comprehensive' and 'software licensing and support'")
    print("- Term changed: one year -> two years")
    print("- Compensation: $10,000 -> $15,000, added annual increase clause")
    print("- Added new Section 5 (Limitation of Liability)")
    print("- Footnotes modified with additional details")


if __name__ == '__main__':
    main()
